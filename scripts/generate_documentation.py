"""Genera la memoria del proyecto en Word, sin incluir credenciales."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'docs' / 'Catalogo_Ferreteria_Documentacion_Integral.docx'
GREEN = '153B31'
LIME = 'D4ED65'
MUTED = '637269'
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2)
sec.right_margin = Cm(2)
sec.header_distance = Cm(.8)
sec.footer_distance = Cm(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(GREEN)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.13
for name, size in [('Title',34),('Subtitle',15),('Heading 1',23),('Heading 2',13),('Heading 3',11)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(GREEN)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.keep_with_next = True
for name in ['List Bullet','List Number']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.size = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(5)
header = sec.header.paragraphs[0]
header.text = 'CATÁLOGO ONLINE  /  DOCUMENTACIÓN DEL PROYECTO'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('Versión 1.0 · 27/10/2026     |     Página ')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(MUTED)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
core = doc.core_properties
core.title = 'Catálogo online para ferreterías · Documentación integral'
core.subject = 'Desarrollo, seguridad, pruebas, publicación y entrega al cliente'
core.author = 'Documentación del proyecto CATALOGOONLINE'
core.keywords = 'ferretería, catálogo, Supabase, Netlify, RLS, MFA, WhatsApp'
core.comments = 'El documento no contiene contraseñas ni claves secretas.'


def p(text='', style=None):
    return doc.add_paragraph(text, style)

def b(text):
    return p(text,'List Bullet')

_step_number = 0
def n(text):
    global _step_number
    _step_number += 1
    para = p(f'{_step_number}. {text}')
    para.paragraph_format.left_indent = Cm(.5)
    para.paragraph_format.first_line_indent = Cm(-.5)
    return para

def h(text):
    global _step_number
    _step_number = 0
    doc.add_heading(text,level=2)

def page(number, title, intro=None):
    doc.add_page_break()
    label=p(f'MEMORIA DEL PROYECTO  /  {number:02d}')
    label.runs[0].font.size=Pt(9)
    label.runs[0].font.bold=True
    label.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    doc.add_heading(title,level=1)
    if intro:p(intro)

def note(title, text):
    t=doc.add_table(rows=1, cols=1)
    cell=t.cell(0,0)
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EDF2E4');tcPr.append(shd)
    cell.paragraphs[0].add_run(title+'\n').bold=True
    cell.paragraphs[0].add_run(text)
    p()

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Cm(width)
    for i,header in enumerate(headers):
        cell=t.rows[0].cells[i]
        cell.text=header
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),GREEN);cell._tc.get_or_add_tcPr().append(shd)
        for run in cell.paragraphs[0].runs:
            run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(9)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cell=cells[i];cell.text=str(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx%2==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F0F3EC');cell._tc.get_or_add_tcPr().append(shd)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(5);para.paragraph_format.space_before=Pt(5)
                for run in para.runs:run.font.size=Pt(9)
        cant=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(cant)
    p()
    return t

def code(text):
    para=p()
    para.paragraph_format.space_after=Pt(9)
    para.paragraph_format.line_spacing=1
    props=para._p.get_or_add_pPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F0F3EC');props.append(shd)
    run=para.add_run(text);run.font.name='Consolas';run.font.size=Pt(8)

def link(label, url):
    para=p()
    hl=OxmlElement('w:hyperlink')
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    hl.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),GREEN);pr.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
    text=OxmlElement('w:t');text.text=label;r.append(text);hl.append(r);para._p.append(hl)

# PORTADA
p('\n\n')
k=p('DESARROLLO · SEGURIDAD · OPERACIÓN');k.runs[0].bold=True;k.runs[0].font.size=Pt(10)
doc.add_heading('Catálogo online\npara ferreterías',0)
doc.add_paragraph('Memoria de desarrollo y guía de entrega del servicio', 'Subtitle')
p('\n')
note('UN CATÁLOGO PÚBLICO. UN PANEL PRIVADO.', 'El negocio administra productos y precios; sus compradores consultan el catálogo y envían el pedido por WhatsApp para coordinar pago y entrega.')
table(['Documento','Referencia'],[
('Versión','1.0 · 27 de octubre de 2026'),
('Proyecto','mateocomar1-spec/CATALOGOONLINE'),
('Tecnología principal','Sitio estático + Supabase + WhatsApp'),
('Alojamiento preparado','Netlify; configuración alternativa para Vercel'),
('Resultados registrados','12 pruebas de lógica, permisos y servidor + 7 de navegador'),
('Estado general','Código publicado en GitHub; publicación web y validación final en producción no confirmadas')],[5,12])
p('Documento dirigido al titular del servicio, a quien implementa la solución y a quien administra la ferretería. Incluye alcance actual, decisiones tomadas, procedimientos y pendientes; no contiene credenciales.')

page(1,'Resumen y estado actual')
p('Se desarrolló una aplicación para que una ferretería publique productos y precios, y reciba consultas o pedidos por WhatsApp. Se agregó un panel privado para administrar el contenido sin modificar el código de la web.')
table(['Componente','Estado al cierre'],[
('Catálogo y panel','Desarrollados y compilados.'),
('Supabase','URL y clave pública configuradas; lectura de tablas reales verificada.'),
('Tablas, RLS y políticas remotas','Confirmadas mediante resultados SQL aportados por el propietario.'),
('Usuario administrador','Membresía confirmada por el propietario mediante SQL.'),
('Pruebas automatizadas','19 aprobadas en las últimas ejecuciones registradas.'),
('Repositorio y paquete','Código y ZIP subidos a la rama main de GitHub.'),
('Netlify y dominio público','Preparados; no se recibió una dirección pública para verificar el despliegue.'),
('Operación real y respaldos','Pendientes de validación y configuración, respectivamente.')],[5,12])
h('Cómo leer el informe')
p('“Implementado” describe lo que está en el código. “Probado localmente” identifica verificaciones automatizadas. “Confirmado por el propietario” corresponde a los resultados SQL compartidos. “Pendiente” no debe presentarse como una funcionalidad o validación ya terminada.')
h('Contenido')
p('02 Recorrido del desarrollo · 03 Arquitectura · 04 Catálogo · 05 Panel y campos · 06 Base de datos · 07 Accesos · 08 Permisos · 09 Seguridad web · 10 Pruebas técnicas · 11 Pruebas de navegador · 12 Evidencia remota · 13 Publicación · 14 Entrega al cliente · 15 Operación y recuperación · 16 Pendientes · 17 Referencias técnicas.')

page(2,'Pasos que desarrollamos')
table(['Etapa','Trabajo realizado'],[
('1. Revisión inicial','Se inspeccionó el repositorio y se comprobó que inicialmente no tenía archivos, commits ni rama main remota.'),
('2. Primer catálogo','Se creó el diseño adaptable a celular, búsqueda, categorías, orden por precio y productos de ejemplo.'),
('3. Pedidos','Se incorporó una bolsa persistente con cantidades, total estimado y mensaje de WhatsApp.'),
('4. Primera integración','Se prepararon Apps Script y plantillas para Google Sheets, según el pedido original.'),
('5. Decisión de arquitectura','Se eligió Supabase por sus permisos, autenticación, validaciones y posibilidad de crecimiento.'),
('6. Panel privado','Se implementaron login, TOTP, edición de productos y datos del negocio e historial de cambios.'),
('7. Base y controles','Se escribieron migración SQL, políticas RLS, funciones de autorización, restricciones y triggers de auditoría.'),
('8. Validación','Se verificaron lógica, roles de base, rutas y formularios en navegador; se corrigieron fallos detectados.'),
('9. Proyecto real','Se configuraron los datos públicos de Supabase y se verificó lectura; el propietario confirmó instalación y autorización por SQL.'),
('10. Publicación preparada','Se agregaron compilación, cabeceras, rutas para Netlify y Vercel, y un ZIP exclusivo de archivos públicos.'),
('11. Entrega por GitHub','Como el enlace local del ZIP no se podía descargar, se subieron el código y el paquete al repositorio elegido.'),
('12. Documentación','Se consolidó el alcance, la seguridad, la evidencia y el procedimiento de entrega en este documento.')],[4,13])
note('Cambio de integración', 'Google Sheets dejó de ser la fuente activa. Sus archivos se conservaron como referencia histórica. La aplicación actual consulta Supabase.')

page(3,'Arquitectura y recorrido de los datos')
p('La web se publica como archivos estáticos. El navegador consulta Supabase mediante su API; las reglas de acceso se aplican en PostgreSQL. No es necesario alojar un servidor propio para la publicación prevista en Netlify.')
table(['Componente','Responsabilidad'],[
('Catálogo público','Muestra productos publicados y datos comerciales. No exige cuenta al comprador.'),
('Panel /admin','Interfaz de trabajo del negocio. Usa autenticación y segundo factor.'),
('Supabase Auth','Gestiona identidad, contraseña, factores TOTP y tokens de sesión.'),
('PostgreSQL / RLS','Controla qué filas y operaciones puede usar cada sesión.'),
('GitHub','Conserva código, documentación y paquete entregable.'),
('Netlify / Vercel','Sirven los archivos compilados por HTTPS y aplican rutas y cabeceras configuradas.'),
('WhatsApp','Recibe la lista solicitada. El personal confirma disponibilidad, pago y entrega.')],[4,13])
h('Flujo del comprador')
code('Visita la web → busca productos → arma su bolsa\n→ envía el mensaje por WhatsApp → coordina con la ferretería')
h('Flujo del administrador')
code('Email y contraseña → cuenta autorizada → código TOTP\n→ guarda cambios → Supabase valida permisos y datos\n→ registra auditoría → el catálogo refleja cambios al recargar')
p('El panel y Supabase son lugares diferentes: /admin es parte de la web del servicio; el dashboard de Supabase es la herramienta técnica del propietario del proyecto. El cliente no necesita el dashboard de Supabase para editar productos.')

page(4,'Catálogo público y pedidos')
h('Funciones desarrolladas')
b('Diseño adaptable a escritorio y celular, con portada, catálogo, instrucciones de compra y contacto.')
b('Búsqueda por nombre o categoría, sin distinguir tildes; filtro por categoría y orden por precio o nombre.')
b('Tarjetas con nombre, presentación, precio, imagen o ícono, destacado y disponibilidad.')
b('Bolsa con cantidades, suma y resta de unidades, eliminación y total estimado.')
b('Persistencia de códigos y cantidades en el navegador; los precios se vuelven a leer de la base al cargar la web.')
b('Envío de un mensaje de WhatsApp con productos, códigos, cantidades y total estimado; consulta general por WhatsApp.')
h('Publicación y disponibilidad')
p('Un producto nuevo comienza como borrador. Solo aparece al público cuando el administrador marca “Publicado”. “Disponible” indica si se puede agregar a la bolsa; no representa una reserva ni un inventario con unidades descontables. Para retirar un producto, se despublica sin borrar su historial.')
h('Qué ocurre ante problemas de conexión')
p('Si falta la configuración o no se puede cargar la base, se muestra un catálogo de demostración identificado y se deshabilita el envío de pedidos. Si la base funciona y no hay productos publicados, la lista queda vacía. No se habilitan compras con precios de ejemplo.')
note('Alcance del pedido', 'El total es estimado. No hay cobro integrado, validación automática de pago, reserva de stock, facturación ni almacenamiento de pedidos en la base. El negocio confirma estos puntos por WhatsApp.')

page(5,'Panel privado: qué puede editar el cliente')
p('El panel se abre desde el enlace “Administrar catálogo”, al pie de la web, o agregando /admin a la dirección publicada. También existe /admin.html. Solo una cuenta autorizada y verificada puede administrar datos.')
table(['Dato / función','Comportamiento actual'],[
('Código','Único e inmutable después de crear el producto.'),
('Nombre y categoría','Editables; permiten organizar y encontrar productos.'),
('Precio de venta','Hasta dos decimales. Validación en formulario y base.'),
('Presentación','Texto breve: unidad, medida, envase u otra característica.'),
('Imagen e ícono','Una URL HTTPS pública opcional y un ícono de respaldo.'),
('Disponible','Permite o bloquea agregar ese producto al pedido.'),
('Destacado','Da prioridad en el orden de destacados.'),
('Publicado','Define si el producto aparece en el catálogo público.'),
('Datos de la ferretería','Nombre, WhatsApp, dirección, horarios, moneda y formato regional.'),
('Historial','Últimos 100 cambios, con fecha, actor y valores anteriores/nuevos.'),
('Actualizar datos','Recarga lista o configuración desde la base.'),
('Cerrar sesión','Limpia la interfaz y descarta el cliente de autenticación en memoria.')],[4,13])
h('Ediciones simultáneas')
p('Se compara updated_at antes de guardar una edición. Si otra persona cambió el mismo registro, la operación se rechaza. Hay que recargar, revisar los valores nuevos y editar otra vez. Así se evita sobrescribir un precio inadvertidamente.')
note('Diferencia con una versión futura', 'Hoy no hay descripción extensa ni subida directa de fotos. Se puede editar nombre y presentación; las fotos se indican por URL. Estos puntos requieren desarrollo adicional si se incluyen en la oferta comercial.')

page(6,'Base de datos e instalación')
table(['Tabla','Contenido y exposición'],[
('public.products','Datos comerciales del producto. El visitante solo lee filas publicadas.'),
('public.store_settings','Datos públicos del negocio. Una fila de configuración con id = 1.'),
('private.catalog_admins','Membresías administrativas vinculadas a usuarios de Auth. No editable desde el panel.'),
('private.catalog_audit','Registro de cambios. El cliente no puede escribir ni borrar directamente su contenido.')],[6,11])
h('Instalación de un proyecto nuevo')
n('Crear un proyecto de Supabase del negocio y conservar el acceso de propietario de forma segura.')
n('En SQL Editor, ejecutar una vez supabase/migrations/001_catalog.sql. La migración es transaccional: ante un error se revierte el conjunto.')
n('Configurar email/contraseña y TOTP, desactivar registros públicos y crear la cuenta administrativa.')
n('Autorizar la cuenta con supabase/authorize-admin.sql, sustituyendo el email del ejemplo.')
n('Completar config.js con Project URL y clave sb_publishable_ del mismo proyecto. Nunca usar una clave secreta o service_role.')
n('Publicar la web y comprobar permisos y operación con una cuenta de prueba.')
h('Situación de este proyecto')
p('Ya se verificó lectura de products y store_settings mediante la API real. El propietario compartió consultas que confirman las cuatro tablas con RLS activado, las políticas esperadas y la función administrativa. También confirmó una fila de autorización de usuario.')
note('No reinstalar sobre datos existentes', 'La migración inicial no es un script para ejecutar repetidamente. Si las tablas ya existen, revisar su esquema y preparar una migración incremental; no borrar tablas para resolver conflictos.')

page(7,'Cuentas, segundo factor y acceso')
h('Ingreso actual')
n('Abrir la dirección pública del panel y escribir el email y la contraseña de la cuenta creada por el propietario.')
n('El sistema comprueba que la cuenta pertenezca a private.catalog_admins. Una cuenta común no puede autorizarse a sí misma.')
n('En la primera sesión, vincular un autenticador TOTP mediante el QR o la clave manual mostrada durante la inscripción.')
n('Ingresar el código de seis dígitos. En los siguientes accesos se vuelve a pedir el código del autenticador.')
n('Solo después de obtener nivel aal2 se permite administrar los datos.')
h('Sesión en memoria')
p('La sesión administrativa no se persiste en localStorage. Recargar la página obliga a ingresar de nuevo. El carrito público guarda únicamente códigos y cantidades. Cerrar sesión descarta el cliente de autenticación incluso si falla la solicitud de red de cierre.')
h('Alta y baja de administradores')
p('Cada persona debe tener su propia cuenta. El propietario crea el usuario en Authentication → Users y agrega su ID a la lista privada mediante el SQL provisto. Para quitar permisos, se elimina su membresía: el control se aplica en la siguiente consulta aunque el token aún conserve aal2.')
code("delete from private.catalog_admins\nwhere user_id = (\n  select id from auth.users\n  where lower(email) = lower('EMAIL_A_REVOCAR')\n);")
note('Entrega y recuperación pendientes de automatizar', 'La aplicación todavía no acepta invitaciones por enlace ni implementa recuperación o cambio de contraseña. No prometer esos flujos. La recuperación actual requiere intervención verificada del propietario con las herramientas administrativas de Supabase; no debe resolverse quitando MFA.')

page(8,'Seguridad de acceso a la base')
p('RLS significa seguridad por fila: PostgreSQL evalúa las políticas antes de devolver o modificar registros. Ocultar botones en el panel no es el control principal; la restricción también opera cuando alguien consulta la API directamente.')
table(['Operación','Visitante','Cuenta común','Admin sin MFA','Admin + MFA'],[
('Ver publicados y datos públicos','Sí','Sí','Sí','Sí'),
('Ver borradores','No','No','No','Sí'),
('Crear / editar productos','No','No','No','Sí'),
('Editar datos del negocio','No','No','No','Sí'),
('Ver historial','No','No','No','Sí'),
('Concederse rol de admin','No','No','No','No'),
('Borrar productos vía cliente','No','No','No','No'),
('Alterar historial directamente','No','No','No','No')],[5.8,2.5,2.6,3,3.1])
h('Condición administrativa comprobada')
code("select coalesce(auth.jwt()->>'aal' = 'aal2', false)\n  and exists (\n    select 1 from private.catalog_admins\n    where user_id = auth.uid()\n  );")
p('La función is_catalog_admin es SECURITY DEFINER y usa search_path vacío con nombres de esquema explícitos. Las membresías se comprueban en una tabla privada, no en un dato editable del perfil. Las funciones tienen permisos de ejecución explícitos.')
h('Otros controles de base')
b('Permisos GRANT/REVOKE restringen operaciones y acceso al esquema privado.')
b('Restricciones validan campos, precios, identificadores y configuración del negocio.')
b('Triggers mantienen la fecha de actualización y generan la auditoría en la misma transacción.')
p('El propietario SQL y las claves privilegiadas pueden eludir RLS. Por eso deben quedar fuera del frontend y del acceso cotidiano del cliente.')

page(9,'Seguridad web y validación de datos')
table(['Control implementado','Protección y alcance'],[
('Clave pública únicamente','El frontend rechaza claves que no tengan formato publishable y restringe la URL al proyecto HTTPS de Supabase.'),
('Renderizado seguro','Los textos de productos y datos se insertan con textContent. El QR se muestra como imagen, sin insertar SVG como HTML ejecutable.'),
('Validación de productos','Código, longitudes, precio explícito y decimales, categoría y URL HTTPS de imagen. La base vuelve a aplicar restricciones.'),
('Datos comerciales permitidos','El formulario y la API de la aplicación usan campos concretos; no guardan secretos ni costos internos.'),
('Content-Security-Policy','Restringe scripts y estilos al sitio; conexiones a Supabase; prohíbe objetos y enmarcado. Permite imágenes HTTPS y data para el QR.'),
('X-Frame-Options / CSP','DENY y frame-ancestors none reducen el riesgo de incluir el panel dentro de otra página.'),
('Nosniff y referer','Evitan interpretación incorrecta de tipos y envío de la dirección de origen al cargar recursos externos.'),
('Permissions-Policy','Deshabilita cámara, micrófono y geolocalización para este sitio.'),
('Archivos públicos separados','Solo se sirven archivos de dist. SQL, servidor, tests, configuración interna y dependencias no van en el paquete desplegado.'),
('Errores y demostración','Una carga fallida no habilita pedidos con precios de prueba. Formularios muestran errores sin imprimir credenciales.')],[5,12])
p('Las cabeceras se probaron en el servidor local y se prepararon para Netlify y Vercel. Su aplicación efectiva debe comprobarse en el dominio final. El acceso HTTPS de producción depende del alojamiento.')
note('Qué significa esta revisión', 'Son controles y pruebas concretos, no una certificación ni un pentest completo. No se realizaron pruebas de intrusión externas, carga masiva, fuerza bruta real ni restauración de backups.')

page(10,'Pruebas técnicas: 12 aprobadas')
p('Última ejecución registrada: npm test. Resultado: 12 aprobadas, 0 fallidas. Incluye cinco casos de permisos ejecutados en PostgreSQL local mediante PGlite, con identidades y niveles JWT simulados.')
table(['N.º','Caso comprobado','Resultado'],[
('1','Búsqueda sin tildes, categorías y orden por precio.','Aprobado'),
('2','Rechazo de precios inválidos e identificadores duplicados al normalizar productos.','Aprobado'),
('3','Número de WhatsApp válido y mensaje con cantidades y total.','Aprobado'),
('4','Visitante ve solo publicados; no puede escribir ni acceder a membresías o auditoría.','Aprobado'),
('5','Cuenta común, incluso con MFA: sin edición, lectura privada ni autoasignación de permisos.','Aprobado'),
('6','Administrador sin segundo factor: modificaciones de productos y negocio bloqueadas.','Aprobado'),
('7','Administrador con MFA: edición y auditoría; precio negativo y cambio de código rechazados; conflicto de actualización detectado.','Aprobado'),
('8','Revocar membresía quita permisos aun con token aal2.','Aprobado'),
('9','Catálogo y panel con CSP; rutas internas no publicadas; método POST al servidor estático rechazado.','Aprobado'),
('10','Configuración rechaza claves secretas y destinos inseguros o no admitidos.','Aprobado'),
('11','Formularios rechazan precio vacío, negativo o no finito, decimales excesivos, código inválido e imagen no HTTPS.','Aprobado'),
('12','Configuración valida WhatsApp y conserva únicamente los campos públicos permitidos.','Aprobado')],[1.1,13.5,2.4])
h('Alcance de la evidencia')
p('Las políticas probadas son las de la migración del repositorio. La autenticación es simulada en estas pruebas SQL: no se verifican la firma de un JWT real ni todos los controles internos del proveedor. Los rechazos se comprueban por permisos, políticas o ausencia de filas actualizadas según el caso.')

page(11,'Pruebas en navegador: 7 aprobadas')
p('Última ejecución registrada: npm run test:browser. Resultado: 7 aprobadas, 0 fallidas. Playwright usa Chromium y respuestas controladas de Supabase para verificar la interfaz y el flujo de uso.')
table(['N.º','Escenario','Qué demuestra'],[
('1','Catálogo conectado y bolsa','Producto real simulado, precio decimal, cantidades tras recargar y URL/mensaje de WhatsApp.'),
('2','Base no disponible','Carga de demostración identificada y botón de pedido deshabilitado.'),
('3','Administración completa','Login, desafío TOTP, edición de precio, datos del negocio, historial, cierre y ausencia de sesión persistida.'),
('4','Primera sesión','Inscripción TOTP antes de editar y limpieza de la clave de inscripción en la interfaz después de verificar.'),
('5','Cuenta no autorizada','No entra al panel aunque la contraseña simulada sea aceptada.'),
('6','Vista móvil','Catálogo y panel sin desbordamiento horizontal en una pantalla de 390 px.'),
('7','Alta de producto','Código propio, precio y creación inicial como borrador.')],[1.1,4.9,11])
h('Verificaciones adicionales registradas')
b('Compilación y arranque del servidor, respuestas HTTP del catálogo y panel y presencia de CSP.')
b('Integridad del ZIP y coincidencia de sus 10 archivos con la última compilación validada.')
b('Recursos JavaScript y CSS referenciados por ambas páginas incluidos en el paquete.')
b('npm audit --omit=dev informó 0 vulnerabilidades conocidas para las dependencias de producción en esa ejecución. No equivale a un análisis completo de seguridad ni cubre dependencias de desarrollo.')
note('Lo que aún exige una prueba real', 'No se confirmó el login TOTP contra el proyecto remoto, la edición desde un sitio publicado, la recepción efectiva de un pedido por WhatsApp ni las cabeceras del dominio final. No se efectuó un pago real.')

page(12,'Comprobaciones del proyecto Supabase real')
h('Lectura desde el entorno')
p('Se realizaron consultas de solo lectura a products y store_settings utilizando la URL y la clave pública configuradas. Ambas respondieron HTTP 200. La política de red inicialmente bloqueó el dominio; se guardó su habilitación y luego se comprobó conectividad.')
h('Evidencia compartida por el propietario')
table(['Consulta / acción','Resultado recibido'],[
('Listado de tablas y RLS','products, store_settings, catalog_admins y catalog_audit con rowsecurity = true.'),
('Listado de políticas','Productos publicados visibles para anon y authenticated. Lectura privada, creación y edición de productos condicionadas por is_catalog_admin. Edición de negocio con la misma condición.'),
('Definición de is_catalog_admin','Exige aal2 y una membresía vinculada a auth.uid(); SECURITY DEFINER y search_path vacío.'),
('Autorización del usuario','El INSERT de membresía devolvió una fila de user_id.')],[5,12])
p('Estos resultados confirman la estructura reportada, pero no reemplazan una prueba de extremo a extremo. El asistente no ejecutó la instalación remota con permisos administrativos: el propietario operó desde SQL Editor y compartió los resultados.')
h('Consultas de revisión, no de reinstalación')
code("select schemaname, tablename, rowsecurity\nfrom pg_tables\nwhere schemaname in ('public','private');\n\nselect schemaname, tablename, policyname, roles,\n       cmd, qual, with_check\nfrom pg_policies\nwhere schemaname in ('public','private');")
p('El documento omite contraseñas, claves secretas, el valor de la clave pública y el identificador personal del administrador. La configuración técnica pública del proyecto está en config.js.')

page(13,'Publicación en Netlify, paso a paso')
p('El código y el ZIP están en GitHub, rama main. Para el titular que ya tiene Netlify, la importación del repositorio evita depender de la descarga del archivo local.')
h('Opción recomendada: importar GitHub')
n('Entrar a Netlify con la cuenta del titular y elegir Add new project → Import an existing project (el nombre puede variar según la interfaz).')
n('Conectar GitHub y conceder acceso al repositorio mateocomar1-spec/CATALOGOONLINE.')
n('Seleccionar la rama main. Usar Node 24, comando npm run build y directorio de publicación dist. netlify.toml contiene esta configuración.')
n('Iniciar Deploy y esperar el resultado. Copiar la dirección HTTPS asignada por Netlify; recién entonces hay una dirección pública verificable.')
n('Abrir el catálogo y el enlace Administrar catálogo. La misma dirección seguida de /admin abre el panel.')
n('Probar login real con TOTP, publicar un producto, modificar su precio, recargar el catálogo y revisar el historial.')
h('Alternativa: usar el ZIP')
p('Desde GitHub abrir artifacts/catalogo-ferreteria.zip y pulsar Download raw. Descomprimirlo y arrastrar a Netlify Drop la carpeta que contiene directamente index.html y admin.html. No subir solo el HTML, ni el repositorio completo.')
h('Rutas y cabeceras')
p('El paquete incluye _redirects para /admin y /admin/, y _headers para CSP y otras medidas. Vercel tiene configuración alternativa en vercel.json. El panel también contiene una indicación noindex; eso no reemplaza la autenticación.')
note('Estado de publicación', 'No se recibió una dirección pública de Netlify en la conversación. No se afirma que la web ya esté desplegada. Publicar el entorno de desarrollo de Codex es distinto de publicar el sitio en Netlify.')

page(14,'Entrega del servicio a cada cliente')
h('Modelo de la versión actual')
p('Cada ferretería debe tener su propio proyecto Supabase y su propia web/configuración. Agregar otro administrador en la base existente le da acceso a la misma ferretería; no crea un negocio aislado. No vender este diseño como una plataforma multicliente compartida sin desarrollar ese aislamiento.')
h('Paquete de entrega recomendado')
b('Dirección pública del catálogo y dirección exacta del panel, con marca y datos del negocio cargados.')
b('Cuenta personal autorizada y guía para configurar el autenticador. Nunca compartir una única cuenta entre todos los clientes.')
b('Instrucciones de productos, precios, publicaciones y fotos por URL; indicar expresamente las funciones aún no incluidas.')
b('Definición escrita de soporte, mantenimiento, titularidad de dominio y cuentas, costos externos y política de respaldos. No se fijaron precios ni planes de proveedores en este desarrollo.')
h('Cómo manejar las credenciales')
p('El objetivo recomendado es una invitación segura para que el cliente defina su contraseña. Ese flujo todavía no está implementado. Hoy el alta es manual; si se usa una credencial inicial, debe entregarse por un canal seguro y sin guardarla en documentos, repositorios o correos permanentes. El cambio de contraseña desde el panel también está pendiente. Resolver este punto antes de ofrecer un alta autoservicio.')
h('Modelo de email de entrega — sin contraseña')
note('Asunto: Acceso al catálogo de [FERRETERÍA]', 'Hola, [NOMBRE]. Tu catálogo está en [URL DEL CATÁLOGO] y el panel en [URL DEL PANEL]. Tu usuario es [EMAIL]. Coordinaremos por un canal seguro el acceso inicial y la configuración del autenticador. Desde el panel podés editar precios, productos y datos del negocio. Las fotos se incorporan por URL. Para soporte: [CONTACTO Y HORARIOS].')

page(15,'Operación, mantenimiento y recuperación')
h('Trabajo cotidiano del negocio')
n('Ingresar al panel con la cuenta propia y el autenticador.')
n('Crear o editar productos. Revisar precio, presentación, disponibilidad e imagen antes de publicar.')
n('Guardar y comprobar el cambio recargando el catálogo. Despublicar si el producto debe retirarse.')
n('Atender WhatsApp y confirmar stock, total, medio de pago y entrega. El pedido no se guarda automáticamente en la base.')
h('Mantenimiento del servicio')
b('Revisar periódicamente accesos administrativos y retirar cuentas que ya no correspondan.')
b('Actualizar dependencias y volver a ejecutar las pruebas antes de desplegar cambios de código.')
b('Comprobar cabeceras, dominio HTTPS, formularios y funcionamiento de la conexión tras cada publicación.')
b('Definir revisión de logs, disponibilidad, errores y cuotas. No se integraron alertas ni monitoreo externo.')
h('Respaldos')
p('Seleccionar y activar la política de backups disponible en el plan real de Supabase, definir frecuencia, retención y objetivos de recuperación. Incluir el esquema private en copias lógicas y contemplar Auth y Storage según corresponda. Mantener copias cifradas fuera del sitio público y probar la restauración en un proyecto separado. No se activaron backups ni se realizó una restauración en esta tarea.')
h('Si se sospecha un acceso indebido')
p('Revocar la membresía afectada, revisar el historial y las sesiones desde las herramientas del proveedor, y preservar la evidencia. Rotar los secretos comprometidos mediante sus procedimientos oficiales, corregir datos tras validación y volver a probar permisos. La auditoría ayuda a investigar; no es una copia de seguridad ni protege frente al propietario SQL malicioso.')

page(16,'Pendientes y lista de lanzamiento')
table(['Prioridad','Acción / alcance pendiente'],[
('Antes del lanzamiento','Confirmar despliegue HTTPS de Netlify, catálogo y /admin; comprobar las cabeceras del dominio final.'),
('Antes del lanzamiento','Probar usuario real, inscripción TOTP, creación/edición, publicación e historial en el proyecto remoto.'),
('Antes del lanzamiento','Revisar registros públicos desactivados, cuentas autorizadas y seguridad de las cuentas propietarias de Supabase, Netlify y GitHub.'),
('Antes del lanzamiento','Cargar negocio, WhatsApp y productos reales; enviar un pedido de prueba y confirmar su recepción.'),
('Antes del lanzamiento','Activar respaldos y validar restauración. Definir soporte y tratamiento de datos personales acorde al uso real.'),
('Mejora de entrega','Desarrollar invitación, definición inicial y cambio de contraseña, recuperación y procedimiento de pérdida de autenticador.'),
('Mejora funcional','Subida directa de fotos con Storage y políticas de acceso; descripción extensa de productos.'),
('Si se ofrece como alcance adicional','Importación CSV/Excel, cantidades de inventario, pedidos guardados, pagos y facturación.'),
('Si se comparte una base entre clientes','Desarrollar negocios/tenants, membresías por negocio y políticas RLS de aislamiento; probar accesos cruzados.'),
('Según necesidad operativa','Monitoreo, protección antiabuso/CAPTCHA y pruebas adicionales de carga o intrusión.')],[4.2,12.8])
h('Cómo describir el servicio comercialmente')
p('Se puede ofrecer como catálogo administrable con pedido por WhatsApp, cuentas autorizadas y segundo factor. No presentarlo como tienda con cobro automático, gestión completa de inventario, aplicación multicliente o sistema invulnerable. Las pruebas registradas respaldan controles específicos, no una garantía absoluta.')

page(17,'Referencias técnicas y archivos entregados')
h('Repositorio y documentos')
link('GitHub · Código del proyecto', 'https://github.com/mateocomar1-spec/CATALOGOONLINE')
link('GitHub · Paquete ZIP para Netlify', 'https://github.com/mateocomar1-spec/CATALOGOONLINE/blob/main/artifacts/catalogo-ferreteria.zip')
p('Referencia del código revisado: commit 7d8ad070572b7daeab265c596fb427f785ecc339. Este Word documenta ese desarrollo y los resultados posteriores compartidos por el propietario. No contiene credenciales.')
table(['Archivos','Propósito'],[
('index.html, app.js, catalog.js, styles.css','Catálogo, bolsa, lógica comercial y diseño.'),
('admin.html, admin.js, admin.css','Panel privado, autenticación, formularios e historial.'),
('supabase-client.js, validation.js, config.js','Cliente de datos, validación y configuración pública.'),
('supabase/migrations/001_catalog.sql','Instalación inicial de tablas, permisos y auditoría.'),
('supabase/authorize-admin.sql','Alta de una membresía administrativa por el propietario.'),
('scripts/build.mjs, server.js','Compilación y servidor para desarrollo/operación propia.'),
('deploy/, netlify.toml, vercel.json','Rutas, cabeceras y configuración de alojamiento.'),
('tests/, catalog.test.js, playwright.config.js','Pruebas técnicas y de navegador.'),
('README.md, PUBLICACION.md','Instrucciones técnicas y publicación guiada.'),
('docs/ y artifacts/','Este documento y el paquete de publicación.')],[7,10])
h('Comandos de referencia')
code('npm ci                 # instalar desde el lockfile\nnpm run dev            # desarrollo con recompilación\nnpm test               # 12 pruebas técnicas\nnpm run test:browser   # 7 pruebas de navegador\nnpm run build          # generar dist/\nnpm start              # compilar e iniciar servidor\nnpm run package:deploy # crear ZIP; requiere Python 3')
p('Solo dist se publica como web. Cambiar precios y datos desde el panel no exige volver a compilar o publicar el sitio; cambiar código o config.js sí. La generación de este Word usa Python y python-docx, fuera de las dependencias necesarias para operar el catálogo.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
print(f'Páginas diseñadas: 18; tablas: {len(doc.tables)}; párrafos: {len(doc.paragraphs)}')
